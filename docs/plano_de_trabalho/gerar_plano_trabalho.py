"""Preenche o modelo oficial FAPEG do Plano de Trabalho e gera a planilha de controle interno.

O modelo (`modelo_fapeg.docx`) é preenchido NO LUGAR: mesmas tabelas, rótulos, colunas e
quebras de página. Só se clonam blocos de linha que o próprio modelo repete (item de despesa,
membro da equipe) e se removem linhas-modelo não usadas. Orientação da tutoria SENAR/FAPEG
(25/09/2026): o modelo não pode ser alterado.

Uso:  python gerar_plano_trabalho.py
PDF:  soffice --headless --convert-to pdf PLANO_DE_TRABALHO_VIA_LEITE_SENSE.docx
"""
import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table

import conteudo as C

AQUI = Path(__file__).parent
MODELO = AQUI / "modelo_fapeg.docx"
SAIDA_DOCX = AQUI / "PLANO_DE_TRABALHO_VIA_LEITE_SENSE.docx"
SAIDA_XLSX = AQUI / "PLANO_DE_TRABALHO_VIA_LEITE_SENSE_Controle_Interno.xlsx"

SUBVENCAO = 60000.00
PARCELA = 30000.00
FONTE = "Arial"

# Índices dos elementos do corpo do modelo (ver inspeção em DECISIONS.md).
EL = dict(cnpj=6, razao=8, endereco=11, contato=14, proponente=19, titulo=26, obj_geral=30,
          oe_a=34, oe_b=37, desc_a=41, desc_b=44, met_a=48, met_b=51, just_a=56, just_b=59,
          res_a=65, res_b=68, eq1_a=72, eq1_b=75, eq2=77, eq3=79, custeio=95, capital=97,
          contra_a=101, contra_b=104)


# --------------------------------------------------------------------------- orçamento
def brl(v):
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def total(it):
    return round(it["qtd"] * it["unit"], 2)


def parcela(it):
    return 1 if max(it["meses"]) <= 6 else 2


def conferir_orcamento():
    for grupo in (C.CAPITAL, C.CUSTEIO, C.CONTRAPARTIDA):
        for it in grupo:
            soma = round(sum(it["meses"].values()), 2)
            assert soma == total(it), f"{it['item']}: meses {soma} != total {total(it)}"
            assert min(it["meses"]) > 6 or max(it["meses"]) <= 6, f"{it['item']} cruza parcelas"
    sub = C.CAPITAL + C.CUSTEIO
    cap, cus = sum(map(total, C.CAPITAL)), sum(map(total, C.CUSTEIO))
    cp = sum(map(total, C.CONTRAPARTIDA))
    p1 = sum(total(i) for i in sub if parcela(i) == 1)
    assert round(cap + cus, 2) == SUBVENCAO, f"subvenção = {cap + cus}"
    assert round(p1, 2) == PARCELA, f"1ª parcela = {p1}"
    assert cp >= 0.10 * SUBVENCAO, f"contrapartida {cp} < 10%"
    for g in (C.CAPITAL, C.CUSTEIO, C.CONTRAPARTIDA):  # ordem de compra
        primeiros = [min(i["meses"]) for i in g]
        assert primeiros == sorted(primeiros), "despesas fora da ordem de compra"
    return cap, cus, cp


def texto_meses(meses):
    ms = sorted(meses)
    if len(ms) == 1:
        return f"no mês {ms[0]}"
    if ms == list(range(ms[0], ms[-1] + 1)):
        return f"nos meses {ms[0]} a {ms[-1]}"
    return "nos meses " + ", ".join(map(str, ms[:-1])) + f" e {ms[-1]}"


def justificativa(it):
    oe = it["oe"]
    vinc = ("Vinculado a todos os objetivos específicos" if oe == "todos" else
            f"Vinculado ao{'s' if ' ' in oe else ''} objetivo{'s' if ' ' in oe else ''} "
            f"específico{'s' if ' ' in oe else ''} {oe}")
    return (f"{it['just']} {vinc}; aquisição {texto_meses(it['meses'])} "
            f"({parcela(it)}ª parcela).")


# --------------------------------------------------------------------------- escrita
def _runs(p, texto, negrito=False, tam=9, italico=False):
    partes, buf, dentro = [], "", False
    for ch in texto:
        if ch == "[" and not dentro:
            if buf:
                partes.append((buf, False))
            buf, dentro = "[", True
        elif ch == "]" and dentro:
            partes.append((buf + "]", True))
            buf, dentro = "", False
        else:
            buf += ch
    if buf:
        partes.append((buf, dentro))
    for trecho, destaque in partes:
        r = p.add_run(trecho)
        r.bold, r.italic = negrito, italico
        r.font.size = Pt(tam)
        r.font.name = FONTE
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
        r._element.rPr.rFonts.set(qn("w:cs"), FONTE)
        if destaque and "PREENCHER" in trecho:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def preencher(cell, blocos, tam=9, alinhamento=None, justificar=False):
    """Substitui o conteúdo da célula mantendo a formatação do 1º parágrafo do modelo.

    blocos: str | lista de str | lista de listas [(texto, negrito), ...] (um parágrafo cada).
    """
    if isinstance(blocos, str):
        blocos = [blocos]
    paras = cell.paragraphs
    base = paras[0]._p
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    for r in list(base):
        if r.tag != qn("w:pPr"):
            base.remove(r)
    modelo = copy.deepcopy(base)
    alvo = paras[0]
    for i, bloco in enumerate(blocos):
        if i:
            novo = copy.deepcopy(modelo)
            alvo._p.addnext(novo)
            alvo = type(paras[0])(novo, paras[0]._parent)
        pf = alvo.paragraph_format
        pf.space_after = Pt(4 if len(blocos) > 1 else 0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        pf.first_line_indent = Pt(0)
        if justificar:
            alvo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif alinhamento is not None:
            alvo.alignment = alinhamento
        trechos = bloco if isinstance(bloco, list) else [(bloco, False)]
        for txt, neg in trechos:
            _runs(alvo, txt, neg, tam)


def caixa(tbl, paragrafos, tam=10, negrito_ate_ponto=False):
    blocos = []
    for p in paragrafos:
        if negrito_ate_ponto and ". " in p[:70]:
            cab, resto = p.split(". ", 1)
            blocos.append([(cab + ". ", True), (resto, False)])
        else:
            blocos.append(p)
    preencher(tbl.rows[0].cells[0], blocos, tam, justificar=True)


def _remover_linhas(tbl, linhas):
    for tr in linhas:
        tbl._tbl.remove(tr)


def tabela_despesas(tbl, itens, linhas_bloco, primeira, ultima, rubrica_fixa=None):
    """Clona o bloco-modelo (linhas `primeira`..`primeira+linhas_bloco-1`) para cada item."""
    trs = tbl._tbl.findall(qn("w:tr"))
    bloco = [copy.deepcopy(tr) for tr in trs[primeira:primeira + linhas_bloco]]
    ancora = trs[primeira - 1]
    originais = trs[primeira:ultima + 1]
    for n, it in enumerate(itens, 1):
        novos = [copy.deepcopy(tr) for tr in bloco]
        for tr in novos:
            ancora.addnext(tr)
            ancora = tr
        linha = Table(tbl._tbl, tbl._parent)
        rows = [r for r in linha.rows if r._tr in novos]
        c = rows[0].cells
        al_c, al_d = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT
        preencher(c[0], str(n), 9, al_c)
        preencher(c[1], [[(it["item"], True)], [(it["espec"], False)]], 8)
        preencher(c[2], rubrica_fixa or it["rubrica"], 8)
        qtd = it["qtd"]
        preencher(c[3], str(int(qtd)) if qtd == int(qtd) else brl(qtd), 9, al_c)
        preencher(c[4], brl(it["unit"]), 9, al_d)
        preencher(c[5], brl(total(it)), 9, al_d)
        texto = justificativa(it)
        if linhas_bloco == 2:  # capital: rótulo e texto na mesma célula
            preencher(rows[1].cells[1], [[("Justificativa: ", False), (texto, False)]], 8,
                      justificar=True)
        else:
            preencher(rows[1].cells[1], "Justificativa:", 8)
            preencher(rows[2].cells[1], texto, 8, justificar=True)
            for par in rows[2].cells[1].paragraphs:
                par.paragraph_format.left_indent = Pt(0)
    _remover_linhas(tbl, originais)


def gerar_docx(cap, cus, cp):
    d = Document(MODELO)
    el = list(d.element.body.iterchildren())
    T = {k: Table(el[i], d._body) for k, i in EL.items()}

    E, P = C.EMPRESA, C.PROPONENTE
    # 1. Empresa / 2. Proponente — valor na linha de cima, rótulo do modelo embaixo
    preencher(T["cnpj"].rows[0].cells[0], E["cnpj"], 10)
    preencher(T["cnpj"].rows[0].cells[2], E["nome_fantasia"], 10)
    preencher(T["razao"].rows[0].cells[0], E["razao_social"], 10)
    preencher(T["endereco"].rows[0].cells[0], E["endereco"], 10)
    preencher(T["contato"].rows[0].cells[0], E["telefone"], 10)
    preencher(T["contato"].rows[0].cells[2], E["email"], 10)
    tp = T["proponente"]
    preencher(tp.rows[0].cells[0], P["cpf"], 10)
    preencher(tp.rows[0].cells[2], P["nome"], 10)
    preencher(tp.rows[2].cells[0], P["telefone"], 10)
    preencher(tp.rows[2].cells[2], P["email"], 10)

    # 3. Dados do projeto
    caixa(T["titulo"], [C.TITULO])
    caixa(T["obj_geral"], [C.OBJETIVO_GERAL])
    slots = [r.cells[0] for r in T["oe_a"].rows] + [r.cells[0] for r in T["oe_b"].rows]
    assert len(C.OBJETIVOS_ESPECIFICOS) <= len(slots)
    for cel, txt in zip(slots, C.OBJETIVOS_ESPECIFICOS):
        preencher(cel, txt, 8.5, justificar=True)
    sobra = len(slots) - len(C.OBJETIVOS_ESPECIFICOS)
    _remover_linhas(T["oe_b"], [r._tr for r in T["oe_b"].rows][len(T["oe_b"].rows) - sobra:])

    caixa(T["desc_a"], C.DESCRICAO[0])
    caixa(T["desc_b"], C.DESCRICAO[1])
    caixa(T["met_a"], C.METODOLOGIA[0], negrito_ate_ponto=False)
    caixa(T["met_b"], C.METODOLOGIA[1])
    caixa(T["just_a"], C.JUSTIFICATIVA[0], negrito_ate_ponto=True)
    caixa(T["just_b"], C.JUSTIFICATIVA[1], negrito_ate_ponto=True)
    n = 1
    for chave, grupo in (("res_a", C.RESULTADOS[0]), ("res_b", C.RESULTADOS[1])):
        blocos = []
        for titulo, texto in grupo:
            blocos.append([(f"{n}. {titulo}: ", True), (texto, False)])
            n += 1
        preencher(T[chave].rows[0].cells[0], blocos, 10, justificar=True)

    # Equipe executora — o modelo traz 3 blocos; os demais repetem o último bloco do modelo
    tabela_eq3 = el[EL["eq3"]]
    ancora = tabela_eq3
    extras = []
    for _ in range(len(C.EQUIPE) - 3):
        sep = copy.deepcopy(el[EL["eq3"] - 1])
        nova = copy.deepcopy(tabela_eq3)
        ancora.addnext(sep)
        sep.addnext(nova)
        ancora = nova
        extras.append(Table(nova, d._body))
    blocos_eq = [(T["eq1_a"], 0, T["eq1_b"], 0, 2)] + [
        (t, 0, t, 2, 4) for t in [T["eq2"], T["eq3"]] + extras]
    for (t_nome, r_nome, t_cont, r_cont, r_form), m in zip(blocos_eq, C.EQUIPE):
        cpf, nome, tel, email, formacao, funcao = m
        preencher(t_nome.rows[r_nome].cells[0], cpf, 10)
        preencher(t_nome.rows[r_nome].cells[2], nome, 10)
        preencher(t_cont.rows[r_cont].cells[0], tel, 10)
        preencher(t_cont.rows[r_cont].cells[2], email, 10)
        preencher(t_cont.rows[r_form].cells[0], formacao, 10)
        preencher(t_cont.rows[r_form].cells[2], funcao, 10)

    # Despesas — blocos do modelo: custeio e contrapartida 3 linhas, capital 2 linhas
    tabela_despesas(T["custeio"], C.CUSTEIO, 3, 3, 26)
    tabela_despesas(T["capital"], C.CAPITAL, 2, 3, 10, rubrica_fixa=C.R_CAPITAL)
    tabela_despesas(T["contra_a"], C.CONTRAPARTIDA, 3, 3, 19)
    tb = T["contra_b"]
    _remover_linhas(tb, [r._tr for r in tb.rows][0:5])  # continuação dos itens 6 e 7 do modelo
    rows = tb.rows
    al_d = WD_ALIGN_PARAGRAPH.RIGHT
    preencher(rows[1].cells[5], brl(cp), 9, al_d)                 # SUB-TOTAL (contrapartida)
    preencher(rows[5].cells[2], brl(cap), 9, al_d)                # CAPITAL
    preencher(rows[6].cells[2], brl(cus), 9, al_d)                # CUSTEIO
    preencher(rows[7].cells[2], brl(cap + cus), 9, al_d)          # TOTAL SUBVENÇÃO
    preencher(rows[8].cells[2], brl(cp), 9, al_d)                 # CONTRAPARTIDA
    preencher(rows[11].cells[1], brl(SUBVENCAO), 9)               # limite
    preencher(rows[11].cells[2], brl(SUBVENCAO - cap - cus), 9)   # saldo

    d.core_properties.title = "Plano de Trabalho — VIA LEITE SENSE Saúde"
    d.core_properties.author = "VIA LEITE SENSE"
    d.save(SAIDA_DOCX)


# --------------------------------------------------------------------------- controle interno
# Evidências de cada objetivo: ficam só no controle interno (orientação da tutoria).
EVIDENCIAS = [
    "Protocolo assinado pela RT veterinária; termos de adesão e consentimento; planilha de "
    "linha de base.",
    "Relatório de ensaio de bancada; certificado de calibração RBC; 30 dias de série da prova "
    "de conceito.",
    "Relatório de testes; registro de versão; demonstração em produção.",
    "Registros de telemetria; relatório mensal de operação.",
    "Laudos do laboratório credenciado; registro de CMT; planilha de eventos e alertas.",
    "Listas de presença; questionários de avaliação; cartilha.",
    "Laudos de tanque antes e depois; cartas de intenção assinadas; relatório final.",
]


def gerar_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    verde = PatternFill("solid", fgColor="1F5C3A")
    claro = PatternFill("solid", fgColor="E3F0E7")
    branco, neg = Font(bold=True, color="FFFFFF"), Font(bold=True)
    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    moeda, wrap = "#,##0.00", Alignment(wrap_text=True, vertical="top")

    def cab(ws, linha, titulos, larg):
        for i, (t, w) in enumerate(zip(titulos, larg), 1):
            c = ws.cell(linha, i, t)
            c.fill, c.font, c.border = verde, branco, borda
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w

    grupos = [("Capital", "Subvenção", "C", C.CAPITAL), ("Custeio", "Subvenção", "K", C.CUSTEIO),
              ("Custeio", "Contrapartida", "P", C.CONTRAPARTIDA)]

    ws = wb.active
    ws.title = "Despesas"
    cab(ws, 1, ["ID", "Fonte", "Natureza", "Parcela", "Rubrica", "Item", "Especificação", "Qtd",
                "Vl unitário", "Total", "Objetivo(s)", "Meses", "Justificativa",
                "Orçamento 1", "Orçamento 2", "Orçamento 3"],
        [6, 13, 10, 8, 34, 40, 55, 8, 12, 13, 11, 14, 60, 14, 14, 14])
    ln = 2
    for nat, fonte, pref, itens in grupos:
        for k, it in enumerate(itens, 1):
            vals = [f"{pref}{k}", fonte, nat, f"{parcela(it)}ª",
                    C.R_CAPITAL if nat == "Capital" else it["rubrica"], it["item"], it["espec"],
                    it["qtd"], it["unit"], f"=H{ln}*I{ln}", it["oe"],
                    ", ".join(f"M{m}" for m in sorted(it["meses"])), justificativa(it)]
            for i, v in enumerate(vals, 1):
                c = ws.cell(ln, i, v)
                c.border, c.alignment = borda, wrap
                if i in (9, 10):
                    c.number_format = moeda
            ln += 1
    ult = ln - 1
    ws.freeze_panes = "A2"

    wd = wb.create_sheet("Desembolso mensal")
    cab(wd, 1, ["ID", "Item", "Fonte"] + [f"M{m}" for m in range(1, 13)] + ["Total", "Confere"],
        [6, 46, 13] + [10] * 12 + [12, 10])
    r = 2
    for nat, fonte, pref, itens in grupos:
        for k, it in enumerate(itens, 1):
            wd.cell(r, 1, f"{pref}{k}")
            wd.cell(r, 2, it["item"])
            wd.cell(r, 3, fonte)
            for m in range(1, 13):
                wd.cell(r, 3 + m, it["meses"].get(m) or None).number_format = moeda
            wd.cell(r, 16, f"=SUM(D{r}:O{r})").number_format = moeda
            wd.cell(r, 17, f'=IF(ROUND(P{r}-SUMIF(Despesas!A:A,A{r},Despesas!J:J),2)=0,"OK",'
                           f'"DIVERGE")')
            r += 1
    ud = r - 1
    for rot, crit in (("Subvenção por mês", "Subvenção"), ("Contrapartida por mês",
                                                           "Contrapartida")):
        wd.cell(r, 2, rot).font = neg
        for col in range(4, 17):
            L = get_column_letter(col)
            c = wd.cell(r, col, f'=SUMIF($C$2:$C${ud},"{crit}",{L}2:{L}{ud})')
            c.number_format, c.font, c.fill = moeda, neg, claro
        r += 1

    wr = wb.create_sheet("Resumo", 0)
    wr["A1"] = "Resumo financeiro — VIA LEITE SENSE Saúde (controle interno)"
    wr["A1"].font = Font(bold=True, size=12)
    for col, w in (("A", 40), ("B", 16), ("C", 36)):
        wr.column_dimensions[col].width = w
    s = f"Despesas!J2:J{ult}"
    f_ = f"Despesas!B2:B{ult}"
    linhas = [
        ("Capital (subvenção)", f'=SUMIFS({s},{f_},"Subvenção",Despesas!C2:C{ult},"Capital")'),
        ("Custeio (subvenção)", f'=SUMIFS({s},{f_},"Subvenção",Despesas!C2:C{ult},"Custeio")'),
        ("TOTAL SUBVENÇÃO", "=B3+B4"),
        ("1ª parcela — meses 1 a 6 (subvenção)",
         f'=SUMIFS({s},{f_},"Subvenção",Despesas!D2:D{ult},"1ª")'),
        ("2ª parcela — meses 7 a 12 (subvenção)",
         f'=SUMIFS({s},{f_},"Subvenção",Despesas!D2:D{ult},"2ª")'),
        ("Contrapartida", f'=SUMIF({f_},"Contrapartida",{s})'),
        ("TOTAL DO PROJETO", "=B5+B8"),
    ]
    for i, (rot, f) in enumerate(linhas, 3):
        wr.cell(i, 1, rot).border = borda
        c = wr.cell(i, 2, f)
        c.number_format, c.border = moeda, borda
        if rot.startswith("TOTAL"):
            wr.cell(i, 1).font = c.font = neg
            wr.cell(i, 1).fill = c.fill = claro
    wr["C5"] = f'=IF(B5={SUBVENCAO},"Utiliza 100% do limite","DIVERGE DO LIMITE")'
    wr["C6"] = f'=IF(B6={PARCELA},"OK — R$ 30 mil","DIVERGE")'
    wr["C7"] = f'=IF(B7={PARCELA},"OK — R$ 30 mil","DIVERGE")'
    wr["C8"] = f'=IF(B8>={SUBVENCAO}*10%,"DENTRO DO LIMITE MÍNIMO EXIGIDO","ABAIXO DE 10%")'

    wo = wb.create_sheet("Objetivos e evidências")
    cab(wo, 1, ["Nº", "Objetivo específico (com prazo)", "Evidência de cumprimento",
                "Status"], [5, 80, 60, 14])
    for i, (oe, ev) in enumerate(zip(C.OBJETIVOS_ESPECIFICOS, EVIDENCIAS), 2):
        for j, v in enumerate([i - 1, oe, ev, "A iniciar"], 1):
            c = wo.cell(i, j, v)
            c.alignment, c.border = wrap, borda
    wb.save(SAIDA_XLSX)


if __name__ == "__main__":
    cap, cus, cp = conferir_orcamento()
    gerar_docx(cap, cus, cp)
    gerar_xlsx()
    print(f"Capital {brl(cap)} | Custeio {brl(cus)} | Subvenção {brl(cap + cus)} | "
          f"1ª parcela {brl(PARCELA)} | Contrapartida {brl(cp)}")
